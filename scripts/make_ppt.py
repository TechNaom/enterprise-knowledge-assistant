#!/usr/bin/env python3
"""Generate the capstone presentation (docs/Enterprise_Knowledge_Assistant.pptx).

    python scripts/make_ppt.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Enterprise_Knowledge_Assistant.pptx"

BLUE = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x1F, 0x2A, 0x44)
GREY = RGBColor(0x55, 0x5F, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _box(slide, x, y, w, h):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    return tf


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def title_slide(title, subtitle, footer):
    s = prs.slides.add_slide(BLANK)
    _bg(s, DARK)
    tf = _box(s, 0.9, 2.4, 11.5, 2)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(0x9E, 0xC5, 0xFF)
    f = _box(s, 0.9, 6.4, 11.5, 0.6)
    fp = f.paragraphs[0]
    fp.text = footer
    fp.font.size = Pt(14)
    fp.font.color.rgb = GREY


def content_slide(title, bullets):
    s = prs.slides.add_slide(BLANK)
    _bg(s, WHITE)
    # title bar
    t = _box(s, 0.7, 0.4, 12, 1)
    tp = t.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(30)
    tp.font.bold = True
    tp.font.color.rgb = BLUE
    # bullets
    body = _box(s, 0.9, 1.6, 11.5, 5.3)
    for i, (text, level) in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(20 if level == 0 else 17)
        p.font.bold = level == 0
        p.font.color.rgb = DARK if level == 0 else GREY
        p.space_after = Pt(6)


def b(text, level=0):
    return (text, level)


# ---- Slides ---------------------------------------------------------------
title_slide(
    "Enterprise Knowledge Assistant",
    "Production-Oriented Advanced RAG over Company Documents",
    "GenAI Development Program — Capstone (Batch 1)  ·  Project 2",
)

content_slide("The Problem", [
    b("Employees waste time hunting through scattered policy documents"),
    b("HR policies, IT security, travel, benefits, code of conduct, FAQs", 1),
    b("A basic 'chat with PDF' is not enough for real use"),
    b("Misses exact terms/acronyms, forgets context, hallucinates, no citations", 1),
    b("Goal: a trustworthy assistant that answers ONLY from company documents,"),
    b("with sources — and remembers the conversation", 1),
])

content_slide("What I Built", [
    b("An Employee Knowledge Assistant powered by advanced RAG"),
    b("Ingests a local folder of PDF + Markdown/TXT documents", 1),
    b("Answers questions with grounded, source-cited responses", 1),
    b("Remembers multi-turn conversations", 1),
    b("Runs fully locally (open-source) or deploys to the cloud"),
    b("Ollama (llama3.2) for local; Google Gemini for the hosted live URL", 1),
    b("Clean, modular Python + Streamlit chat interface"),
])

content_slide("Advanced Retrieval Pipeline — beyond chat-with-PDF", [
    b("1. Multi-query expansion — LLM rephrases the question to widen recall"),
    b("2. Hybrid retrieval — dense vectors (ChromaDB) + BM25 keyword search"),
    b("Catches meaning AND exact terms (PTO, VPN, policy IDs)", 1),
    b("3. Reciprocal Rank Fusion — robustly merges both rankings"),
    b("4. LLM re-ranking — one batched call scores candidates for relevance"),
    b("5. MMR diversification — relevant AND non-redundant final context"),
    b("6. History-aware condensing — resolves follow-ups across turns"),
])

content_slide("Architecture", [
    b("Ingestion (offline):"),
    b("load PDF/MD/TXT -> chunk (overlap) -> embed -> ChromaDB (dense) + BM25 (sparse)", 1),
    b("Query (online):"),
    b("condense -> multi-query -> dense + sparse -> RRF -> re-rank -> MMR -> top-k", 1),
    b("grounded answer with [n] citations", 1),
    b("Provider abstraction:"),
    b("one interface, swap Ollama <-> Gemini via a single config value", 1),
])

content_slide("Tech Stack", [
    b("Language: Python 3"),
    b("RAG: ChromaDB (vector DB) · BM25 · Reciprocal Rank Fusion · MMR"),
    b("Models: Ollama (llama3.2, nomic-embed-text) · Google Gemini"),
    b("gemini-flash-lite-latest + gemini-embedding-001", 1),
    b("Interface: Streamlit chat UI (+ CLI)"),
    b("Documents: pypdf (PDF), native Markdown/TXT"),
    b("No paid infrastructure required; secrets kept out of source (.env)"),
])

content_slide("Key Features & Results", [
    b("Grounded & faithful — answers only from context, declines when unknown"),
    b("Source citations on every answer [1], [2] ..."),
    b("Multi-turn memory — 'And how long for partners?' resolves correctly"),
    b("Two document formats supported (PDF + Markdown/TXT)"),
    b("Fast on Gemini (~1-10s/query); fully offline option on Ollama"),
    b("Validated end-to-end on a fictional company (Aveltra Technologies)", 1),
])

content_slide("Demo", [
    b("Live app: Streamlit chat interface"),
    b("Ask: 'How many paid vacation days and can I carry them over?'", 1),
    b("-> 20 days PTO, up to 5 carried over  [leave-policy]", 1),
    b("Ask a follow-up to show conversation memory"),
    b("Ask an out-of-scope question -> grounded refusal"),
    b("Expand 'Sources' and 'Retrieval trace' to show the pipeline working"),
])

content_slide("Summary", [
    b("Delivered a production-oriented, advanced-RAG knowledge assistant"),
    b("Advanced retrieval, conversation memory, grounded cited answers, clean UI"),
    b("Local-first and open-source, with a hosted live deployment"),
    b("Modular, documented, reproducible (README + sample data + sample output)"),
    b("GitHub: github.com/TechNaom/enterprise-knowledge-assistant"),
])

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"Wrote {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
